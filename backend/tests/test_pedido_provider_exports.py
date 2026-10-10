import os
import unittest
import zipfile
from datetime import date
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

os.environ.setdefault("SECRET_KEY", "clave_de_pruebas_no_usar_en_produccion")

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config.database import get_db
from app.middlewares.auth_middleware import get_current_user
from app.models.user_model import UserRole
from app.routes.pedido_routes import router
from app.services import pedido_service


class ProviderExportTests(unittest.TestCase):
    def setUp(self):
        self.provider = {
            "proveedor_id": 1,
            "proveedor_nombre": "Proveedor prueba",
            "localidad_id": 1,
            "localidad_nombre": "Localidad prueba",
            "ingredientes": [{
                "ingrediente_id": 1,
                "ingrediente_nombre": "Aceite",
                "unidad": "unidades",
                "contenido_por_unidad": "900",
                "unidad_contenido": "cc",
                "cantidad_total": "3.00",
                "cantidad_contenido_total": "2700.00",
                "precio_unitario": "123456.79",
                "costo_total": "370370.37",
                "escuelas": [
                    {"escuela_id": 1, "escuela_codigo": "E1", "escuela_nombre": "Escuela uno", "cantidad": "2.00", "cantidad_contenido": "1800.00"},
                    {"escuela_id": 2, "escuela_codigo": "E2", "escuela_nombre": "Escuela dos", "cantidad": "1.00", "cantidad_contenido": "900.00"},
                ],
            }],
        }
        self.snapshot = {
            "semana_inicio": "2026-10-12",
            "proveedores": [self.provider],
            "escuelas": [
                {"escuela_id": 1, "localidad_id": 1, "ingredientes": []},
                {"escuela_id": 2, "localidad_id": 1, "ingredientes": []},
            ],
        }
        self.pedido = SimpleNamespace(id=1, semana_inicio=date(2026, 10, 12), datos_snapshot=self.snapshot)
        self.user = SimpleNamespace(role=UserRole.admin)

    def test_excel_without_prices_keeps_quantities_and_contains_no_money(self):
        content = pedido_service._provider_excel(self.snapshot, self.provider, include_prices=False)
        workbook = load_workbook(content)
        self.assertEqual(workbook.sheetnames, ["Orden proveedor"])
        rows = list(workbook.active.values)
        self.assertEqual(rows[5], ("Ingrediente", "Unidad", "E1 - Escuela uno", "E2 - Escuela dos", "TOTAL"))
        self.assertEqual(rows[6], ("Aceite", "unidades (900 cc c/u)", "2.00", "1.00", "3.00 (2.7 litros)"))
        self.assertEqual(len(rows), 7)
        for sheet in workbook:
            for row in sheet:
                for cell in row:
                    value = str(cell.value).lower()
                    for forbidden in ("precio", "costo", "123456.79", "370370.37", "$"):
                        self.assertNotIn(forbidden, value)

    def test_excel_default_retains_prices_and_total(self):
        rows = list(load_workbook(pedido_service._provider_excel(self.snapshot, self.provider)).active.values)
        self.assertEqual(rows[5][-2:], ("Precio unit.", "Costo estimado"))
        self.assertEqual(rows[6][-2:], ("123456.79", "370370.37"))
        self.assertEqual(rows[7][-1], "370370.37")

    def test_pdf_without_prices_removes_cost_and_retains_order_details(self):
        # Uncompressed page streams let us inspect the actual PDF text without
        # introducing a PDF reader dependency into the application's test suite.
        with patch("reportlab.rl_config.pageCompression", 0):
            priced = pedido_service._provider_pdf(self.snapshot, self.provider).getvalue()
            unpriced = pedido_service._provider_pdf(self.snapshot, self.provider, include_prices=False).getvalue()
        self.assertIn(b"Costo estimado total: 370370.37", priced)
        for forbidden in (b"Costo", b"Precio", b"123456.79", b"370370.37"):
            self.assertNotIn(forbidden, unpriced)
        for detail in (b"ORDEN DE COMPRA", b"2026-10-12", b"Proveedor prueba", b"Localidad prueba", b"Aceite", b"E1", b"E2", b"3.00", b"2.7 litros"):
            self.assertIn(detail, unpriced)

    def test_routes_export_both_versions_and_apply_school_filter(self):
        app = FastAPI()
        app.include_router(router)
        app.dependency_overrides[get_db] = lambda: None
        app.dependency_overrides[get_current_user] = lambda: self.user
        with TestClient(app) as client, patch.object(pedido_service, "get_pedido_for_user", return_value=self.pedido), patch("reportlab.rl_config.pageCompression", 0):
            for file_format, extension in (("pdf", "pdf"), ("excel", "xlsx")):
                for include_prices in (True, False):
                    with self.subTest(format=file_format, include_prices=include_prices):
                        params = {"localidad_id": 1, "proveedor_id": 1, "escuela_id": 1}
                        # Existing callers omit the flag and still receive prices.
                        if not include_prices:
                            params["incluir_precios"] = "false"
                        response = client.get(f"/pedidos/1/export/proveedores/{file_format}", params=params)
                        self.assertEqual(response.status_code, 200)
                        self.assertEqual(response.headers["content-type"], "application/zip")
                        self.assertEqual("_sin_precios" in response.headers["content-disposition"], not include_prices)
                        with zipfile.ZipFile(BytesIO(response.content)) as archive:
                            names = archive.namelist()
                            self.assertEqual(len(names), 1)
                            self.assertTrue(names[0].endswith(f".{extension}"))
                            self.assertEqual("_sin_precios" in names[0], not include_prices)
                            content = archive.read(names[0])
                        if file_format == "excel":
                            rows = list(load_workbook(BytesIO(content)).active.values)
                            self.assertEqual(rows[5][2:4], ("E1 - Escuela uno", "TOTAL"))
                            self.assertEqual(rows[6][2:4], ("2.00", "2.00 (1.8 litros)"))
                            self.assertEqual("Precio unit." in rows[5], include_prices)
                            self.assertEqual(len(rows), 8 if include_prices else 7)
                        else:
                            self.assertIn(b"E1", content)
                            self.assertNotIn(b"E2", content)
                            self.assertIn(b"1.8 litros", content)
                            self.assertEqual(b"Costo estimado total: 246913.58" in content, include_prices)


if __name__ == "__main__":
    unittest.main()
